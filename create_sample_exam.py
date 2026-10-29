# -*- coding: utf-8 -*-
"""
Script tạo Đề thi Mẫu chuẩn cấu trúc GDPT 2018 (3 Phần: TNKQ 4 lựa chọn, Đúng/Sai, Trả lời ngắn)
có nhúng cả Công thức Toán học OMML (<m:oMath>) thực tế, chỉ số trên/dưới và Hình ảnh đồ thị hàm số
để kiểm thử toàn diện bộ trộn đề `shuffler.py` cũng như cung cấp file mẫu cho người dùng tải về.
"""

import io
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np


def make_omml_fraction(num_text: str, den_text: str):
    """Tạo phần tử công thức phân số Word Equation (<m:oMath>) thực sự."""
    xml = f"""
    <m:oMath {nsdecls('m', 'w')}>
      <m:f>
        <m:num><m:r><m:t>{num_text}</m:t></m:r></m:num>
        <m:den><m:r><m:t>{den_text}</m:t></m:r></m:den>
      </m:f>
    </m:oMath>
    """
    return parse_xml(xml)


def make_omml_integral(lower: str, upper: str, expr: str):
    """Tạo phần tử công thức tích phân Word Equation (<m:oMath>) thực sự."""
    xml = f"""
    <m:oMath {nsdecls('m', 'w')}>
      <m:nary>
        <m:naryPr>
          <m:chr m:val="∫"/>
          <m:limLoc m:val="subSup"/>
        </m:naryPr>
        <m:sub><m:r><m:t>{lower}</m:t></m:r></m:sub>
        <m:sup><m:r><m:t>{upper}</m:t></m:r></m:sup>
        <m:e><m:r><m:t>{expr}</m:t></m:r></m:e>
      </m:nary>
    </m:oMath>
    """
    return parse_xml(xml)


def make_graph_image_bytes() -> bytes:
    """Vẽ một đồ thị hàm số bậc ba y = x^3 - 3x + 1 làm hình minh họa nhúng vào đề Word."""
    fig, ax = plt.subplots(figsize=(3.4, 2.3), dpi=150)
    x = np.linspace(-2.1, 2.1, 200)
    y = x**3 - 3 * x + 1
    ax.plot(x, y, color="#1D4ED8", linewidth=2, label=r"$y = f(x)$")
    ax.axhline(0, color="#334155", linewidth=0.9)
    ax.axvline(0, color="#334155", linewidth=0.9)
    ax.scatter([-1, 1], [3, -1], color="#DC2626", zorder=5, s=25)
    ax.text(-1.15, 3.25, "(-1; 3)", fontsize=8, color="#DC2626")
    ax.text(0.85, -1.55, "(1; -1)", fontsize=8, color="#DC2626")
    ax.set_xlim(-2.3, 2.3)
    ax.set_ylim(-2.5, 4.2)
    ax.grid(True, linestyle="--", alpha=0.4)
    ax.set_title("Đồ thị hàm số y = f(x)", fontsize=9, fontweight="bold")
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    plt.close(fig)
    return buf.getvalue()


def add_option_line(doc, letter: str, text_after: str, is_correct: bool = False, omml_elem=None):
    """Thêm dòng phương án A/B/C/D hoặc a/b/c/d với chữ cái đầu được gạch chân nếu là đáp án đúng."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r_lbl = p.add_run(letter)
    r_lbl.bold = True
    r_lbl.underline = is_correct
    r_lbl.font.name = "Times New Roman"
    r_lbl.font.size = Pt(12)

    sep = ". " if letter.isupper() else ") "
    r_sep = p.add_run(sep + text_after)
    r_sep.font.name = "Times New Roman"
    r_sep.font.size = Pt(12)

    if omml_elem is not None:
        p._p.append(omml_elem)
    return p


def build_sample_exam_docx(target_path: str) -> bytes:
    doc = docx.Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("ĐỀ GỐC MẪU CHUẨN CẤU TRÚC MỚI (3 PHẦN - CÓ CÔNG THỨC & HÌNH ẢNH)")
    r_t.bold = True
    r_t.font.size = Pt(13)

    # =========================================================================
    # PHẦN I: TRẮC NGHIỆM 4 LỰA CHỌN
    # =========================================================================
    p_p1 = doc.add_paragraph()
    r_p1 = p_p1.add_run("PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn.")
    r_p1.bold = True

    # Câu 1 (Có công thức phân số OMML và gạch chân đáp án đúng)
    p_q1 = doc.add_paragraph()
    r_q1 = p_q1.add_run("Câu 1: Cho hàm số phân thức ")
    r_q1.bold = True
    p_q1._p.append(make_omml_fraction("2x + 1", "x - 1"))
    p_q1.add_run(". Tiệm cận ngang của đồ thị hàm số đã cho là đường thẳng nào dưới đây?")
    add_option_line(doc, "A", "y = 2.", is_correct=True)
    add_option_line(doc, "B", "y = 1.", is_correct=False)
    add_option_line(doc, "C", "x = 1.", is_correct=False)
    add_option_line(doc, "D", "y = -1.", is_correct=False)

    # Câu 2 (Có hình ảnh đồ thị nhúng trực tiếp trong thân câu hỏi)
    p_q2 = doc.add_paragraph()
    p_q2.add_run("Câu 2: ").bold = True
    p_q2.add_run("Cho hàm số bậc ba y = f(x) có đồ thị như hình vẽ bên dưới. Giá trị cực đại của hàm số đã cho bằng bao nhiêu?")
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_img = p_img.add_run()
    r_img.add_picture(io.BytesIO(make_graph_image_bytes()), width=Inches(2.8))
    add_option_line(doc, "A", "x = -1.", is_correct=False)
    add_option_line(doc, "B", "y = 3.", is_correct=True)
    add_option_line(doc, "C", "y = -1.", is_correct=False)
    add_option_line(doc, "D", "x = 1.", is_correct=False)

    # Câu 3 (Có công thức tích phân OMML ở cả câu hỏi và phương án)
    p_q3 = doc.add_paragraph()
    p_q3.add_run("Câu 3: ").bold = True
    p_q3.add_run("Giá trị của tích phân xác định ")
    p_q3._p.append(make_omml_integral("0", "2", "3x² dx"))
    p_q3.add_run(" bằng:")
    add_option_line(doc, "A", "8.", is_correct=True)
    add_option_line(doc, "B", "6.", is_correct=False)
    add_option_line(doc, "C", "4.", is_correct=False)
    add_option_line(doc, "D", "12.", is_correct=False)

    # Câu 4 (Hóa học / Vật lý có chỉ số trên và chỉ số dưới H2SO4, Fe3+)
    p_q4 = doc.add_paragraph()
    p_q4.add_run("Câu 4: ").bold = True
    p_q4.add_run("Trong dung dịch loãng, phương trình phân li của axit sunfuric H")
    r_sub2 = p_q4.add_run("2")
    r_sub2.font.subscript = True
    p_q4.add_run("SO")
    r_sub4 = p_q4.add_run("4")
    r_sub4.font.subscript = True
    p_q4.add_run(" tạo ra ion sunfat có kí hiệu hóa học là:")
    add_option_line(doc, "A", "SO3(2-).", is_correct=False)
    add_option_line(doc, "B", "HSO3(-).", is_correct=False)
    p_opt_c = add_option_line(doc, "C", "SO", is_correct=True)
    rc1 = p_opt_c.add_run("4")
    rc1.font.subscript = True
    rc2 = p_opt_c.add_run("2-")
    rc2.font.superscript = True
    p_opt_c.add_run(".")
    add_option_line(doc, "D", "S(2-).", is_correct=False)

    # =========================================================================
    # PHẦN II: TRẮC NGHIỆM ĐÚNG / SAI (a, b, c, d)
    # =========================================================================
    p_p2 = doc.add_paragraph()
    r_p2 = p_p2.add_run("PHẦN II. Câu trắc nghiệm đúng sai.")
    r_p2.bold = True

    # Câu 1 Phần II (Gạch chân các ý Đúng: a và c đúng, b và d sai)
    p_p2_q1 = doc.add_paragraph()
    p_p2_q1.add_run("Câu 1: ").bold = True
    p_p2_q1.add_run("Cho hàm số y = f(x) = x³ - 3x + 1 xác định và liên tục trên R. Xét tính đúng sai của các mệnh đề sau:")
    add_option_line(doc, "a", "Đạo hàm của hàm số đã cho là f'(x) = 3x² - 3.", is_correct=True)
    add_option_line(doc, "b", "Hàm số đồng biến trên khoảng (-1; 1).", is_correct=False)
    add_option_line(doc, "c", "Điểm cực đại của đồ thị hàm số có tọa độ là (-1; 3).", is_correct=True)
    add_option_line(doc, "d", "Giá trị lớn nhất của hàm số trên đoạn [0; 2] bằng 1.", is_correct=False)

    # Câu 2 Phần II (Có công thức phân số OMML trong mệnh đề)
    p_p2_q2 = doc.add_paragraph()
    p_p2_q2.add_run("Câu 2: ").bold = True
    p_p2_q2.add_run("Một vật dao động điều hòa dọc theo trục Ox với phương trình x = 4cos(2πt) (cm). Xét các nhận định sau:")
    add_option_line(doc, "a", "Biên độ dao động của vật bằng 4 cm.", is_correct=True)
    add_option_line(doc, "b", "Chu kì dao động của vật bằng 2 giây.", is_correct=False)
    p_opt_2c = add_option_line(doc, "c", "Tần số dao động của vật được tính bởi công thức ", is_correct=True, omml_elem=make_omml_fraction("ω", "2π"))
    p_opt_2c.add_run(" và bằng 1 Hz.")
    add_option_line(doc, "d", "Tại thời điểm t = 0,5 giây, vật đang ở vị trí biên âm x = -4 cm.", is_correct=True)

    # =========================================================================
    # PHẦN III: TRẮC NGHIỆM TRẢ LỜI NGẮN
    # =========================================================================
    p_p3 = doc.add_paragraph()
    r_p3 = p_p3.add_run("PHẦN III. Câu trắc nghiệm trả lời ngắn.")
    r_p3.bold = True

    p_p3_q1 = doc.add_paragraph()
    p_p3_q1.add_run("Câu 1: ").bold = True
    p_p3_q1.add_run("Một doanh nghiệp sản xuất mỗi ngày được x sản phẩm (1 ≤ x ≤ 100). Tổng lợi nhuận thu được (đơn vị: triệu đồng) được mô tả bởi hàm số P(x) = -x² + 80x - 700. Hỏi doanh nghiệp cần sản xuất bao nhiêu sản phẩm mỗi ngày để đạt lợi nhuận lớn nhất?")
    p_ans1 = doc.add_paragraph("Đáp án: 40")

    p_p3_q2 = doc.add_paragraph()
    p_p3_q2.add_run("Câu 2: ").bold = True
    p_p3_q2.add_run("Cho hình chóp S.ABC có đáy ABC là tam giác vuông cân tại B, AB = BC = 3a, cạnh bên SA vuông góc với mặt phẳng đáy và SA = 4a. Khi a = 1, thể tích khối chóp S.ABC bằng bao nhiêu?")
    p_ans2 = doc.add_paragraph("Đáp án: 6")

    p_end = doc.add_paragraph("-------------- HẾT --------------")
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER

    os.makedirs(os.path.dirname(target_path), exist_ok=True)
    doc.save(target_path)
    with open(target_path, "rb") as f:
        return f.read()


if __name__ == "__main__":
    sample_path = "/working_dir/c_c6db43316bef118a/artifacts/file_generation/ttl=63d/intermediate/tron_de_online/static/de_mau_chuan_2025.docx"
    data = build_sample_exam_docx(sample_path)
    print(f"Created sample exam at {sample_path} ({len(data)} bytes)")
